#!/usr/bin/env python3
"""Gera a planilha de arquitetura de SEO a partir do repositorio.

    python3 bin/seo-planilha.py [destino.xlsx]

Padrao: ~/Desktop/Arquitetura SEO Ajuda Atletis.xlsx

E a planilha que vai para quem cuida de SEO. Regenere sempre que mexer em
estrutura, title ou description, para ela nao divergir do site.

Precisa de openpyxl:  pip install openpyxl
"""
import json, os, re, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit('openpyxl nao instalado. Rode: pip install openpyxl')

DOM = 'https://ajuda.atletis.com.br'
ARIAL = 'Arial'

def dados():
    d = json.load(open('docs.json', encoding='utf-8'))
    suf, nav, out = f" - {d['name']}", d['navigation'], []
    for lang in nav.get('languages', [{'tabs': nav.get('tabs', [])}]):
        for tab in lang.get('tabs', []):
            def add(slug, grupo):
                s = open(f'{slug}.mdx', encoding='utf-8').read()
                fm = re.match(r'^---\n(.*?)\n---\n', s, re.S).group(1)
                g = lambda k: (re.search(rf'^{k}:\s*"(.*?)"\s*$', fm, re.M) or [None, ''])[1]
                kw = re.search(r'^keywords:\s*(\[.*?\])\s*$', fm, re.M | re.S)
                corpo = s[s.index('---', 3) + 3:]
                out.append({'url': DOM + ('/' if slug == 'index' else f'/{slug}'),
                            'modulo': tab.get('tab', ''), 'grupo': grupo or '—',
                            'title': g('title'), 'render': g('title') + suf,
                            'desc': g('description'),
                            'kw': ', '.join(json.loads(kw.group(1))) if kw else '',
                            'palavras': len(re.sub(r'<[^>]+>', '', corpo).split())})
            for pg in tab.get('pages', []): add(pg, '')
            for gr in tab.get('groups', []):
                for pg in gr.get('pages', []): add(pg, gr['group'])
    return out, suf

def main():
    rows, suf = dados()
    destino = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser(
        '~/Desktop/Arquitetura SEO Ajuda Atletis.xlsx')
    HDR = PatternFill('solid', fgColor='FF6F00')
    HDRF = Font(name=ARIAL, bold=True, color='FFFFFF', size=10)
    BODY, BOLD = Font(name=ARIAL, size=10), Font(name=ARIAL, size=10, bold=True)
    WRAP = Alignment(wrap_text=True, vertical='top')
    TOP = Alignment(vertical='top')
    THIN = Border(*[Side(style='thin', color='DDDDDD')] * 4)
    YEL = PatternFill('solid', fgColor='FFF3CD')
    wb = Workbook()

    ws = wb.active; ws.title = 'Leia-me'
    ws.column_dimensions['A'].width = 4; ws.column_dimensions['B'].width = 118
    ws['B1'] = 'Arquitetura de SEO — ajuda.atletis.com.br'
    ws['B1'].font = Font(name=ARIAL, bold=True, size=14)
    def line(t, bold=False, fill=None):
        c = ws.cell(ws.max_row + 1, 2, t)
        c.font = BOLD if bold else BODY; c.alignment = WRAP
        if fill: c.fill = fill
    from datetime import date
    line(f'Gerado por bin/seo-planilha.py a partir do repositorio Atletis-LTDA/Mint em '
         f'{date.today().strftime("%d/%m/%Y")}. Regenere depois de mexer em estrutura.')
    line('')
    line('Abas', bold=True)
    line('Páginas — as URLs no ar, com módulo, grupo, title, meta description, keywords e volume de texto.')
    line('Árvore — a hierarquia da navegação.')
    line('')
    line('Como o title funciona aqui', bold=True)
    line(f'O Mintlify acrescenta "{suf}" ao title de toda página. A coluna "Title renderizado" é o que '
         'sai no HTML e conta na SERP — é ele que precisa ficar entre 50 e 60 caracteres.')
    line('As contagens são um retrato da geração, gravadas como valor. Editando aqui, recalcule à mão.')
    line('')
    line('Ponto em aberto', bold=True, fill=YEL)
    line('A home deste subdomínio precisa de target próprio (ajuda, central de ajuda, como funciona) '
         'em vez de disputar a mesma intenção de busca do site principal atletis.com.br. '
         'A mesma pergunta vale para a porta de cada módulo.', fill=YEL)
    line('')
    line('Estado técnico', bold=True)
    line('Nenhuma página tem noindex. Canonical de cada página aponta para si mesma. '
         'Sitemap em /sitemap.xml e índice de páginas em /llms.txt.')
    line('Navegação configurada em pt-BR, mas o Mintlify emite <html lang="en"> fixo — '
         'limitação da plataforma, ver SEO.md.')

    ws = wb.create_sheet('Páginas')
    cols = [('URL', 46), ('Módulo', 19), ('Grupo', 22), ('Title (no arquivo)', 40),
            ('Title renderizado', 46), ('Nº', 5), ('Meta description', 62), ('Nº', 5),
            ('Keywords', 44), ('Palavras no corpo', 10)]
    for i, (h, w) in enumerate(cols, 1):
        c = ws.cell(1, i, h); c.font = HDRF; c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    for n, r in enumerate(rows, 2):
        for i, v in enumerate([r['url'], r['modulo'], r['grupo'], r['title'], r['render'],
                               len(r['render']), r['desc'], len(r['desc']), r['kw'],
                               r['palavras']], 1):
            c = ws.cell(n, i, v); c.font = BODY; c.border = THIN
            c.alignment = WRAP if i in (1, 4, 5, 7, 9) else TOP
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:J{len(rows)+1}'

    ws = wb.create_sheet('Árvore')
    ws.column_dimensions['A'].width = 4; ws.column_dimensions['B'].width = 58
    ws.column_dimensions['C'].width = 46
    ws['B1'] = 'Hierarquia atual'; ws['B1'].font = Font(name=ARIAL, bold=True, size=14)
    ws['B2'], ws['C2'] = 'Estrutura', 'URL'
    for cel in ('B2', 'C2'): ws[cel].font = HDRF; ws[cel].fill = HDR
    arvore, r = {}, 3
    for x in rows: arvore.setdefault(x['modulo'], {}).setdefault(x['grupo'], []).append(x)
    for mod, grupos in arvore.items():
        ws.cell(r, 2, f'{mod}  ({sum(len(v) for v in grupos.values())} páginas)').font = \
            Font(name=ARIAL, bold=True, size=11); r += 1
        for g, pgs in grupos.items():
            if g != '—':
                ws.cell(r, 2, f'    {g}').font = Font(name=ARIAL, bold=True, size=10, color='666666'); r += 1
            for p in pgs:
                ws.cell(r, 2, f'        {p["title"]}').font = BODY
                ws.cell(r, 3, p['url'].replace(DOM, '') or '/').font = BODY
                r += 1
        r += 1

    os.makedirs(os.path.dirname(destino) or '.', exist_ok=True)
    wb.save(destino)
    print(f'{len(rows)} páginas -> {destino}')

if __name__ == '__main__':
    sys.exit(main())
